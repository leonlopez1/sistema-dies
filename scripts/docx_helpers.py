"""Helpers reutilizables para todos los documentos DOCX del sistema."""

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AZUL = RGBColor(0x1A, 0x2B, 0x4A)
DORADO = RGBColor(0xC9, 0x97, 0x3A)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
GRIS_BG = RGBColor(0xF5, 0xF5, 0xF5)
GRIS_TEXTO = RGBColor(0x77, 0x77, 0x77)
NEGRO = RGBColor(0x22, 0x22, 0x22)
ROJO = RGBColor(0xC0, 0x39, 0x2B)
VERDE = RGBColor(0x27, 0xAE, 0x60)
AMARILLO = RGBColor(0xF3, 0x9C, 0x12)


def set_cell_bg(cell, color_hex: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)


def new_doc() -> Document:
    doc = Document()
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(11)
    return doc


def add_header_footer(doc, title: str, version: str = "v1.0 | Marzo 2026"):
    section = doc.sections[0]
    # Header
    hdr = section.header
    hdr.is_linked_to_previous = False
    p = hdr.paragraphs[0]
    p.clear()
    r1 = p.add_run(f"SISTEMA COMERCIAL OPERATIVO | León López")
    r1.font.size = Pt(9)
    r1.font.bold = True
    r1.font.color.rgb = AZUL
    r2 = p.add_run(f"{'  ' * 30}[LOGO]")
    r2.font.size = Pt(9)
    r2.font.color.rgb = DORADO

    # Footer
    ftr = section.footer
    ftr.is_linked_to_previous = False
    pf = ftr.paragraphs[0]
    pf.clear()
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rf = pf.add_run(f"{title} | {version} | Confidencial")
    rf.font.size = Pt(8)
    rf.font.color.rgb = GRIS_TEXTO


def portada(doc, titulo: str, subtitulo: str, descripcion: str = ""):
    """Crea una portada estándar."""
    p_logo = doc.add_paragraph()
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(30)
    r = p_logo.add_run("[LOGO DE LA EMPRESA]")
    r.font.size = Pt(11)
    r.font.color.rgb = GRIS_TEXTO
    r.font.italic = True

    doc.add_paragraph()

    # Title block
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_bg(cell, "1A2B4A")

    p_t = cell.paragraphs[0]
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t.paragraph_format.space_before = Pt(24)
    p_t.paragraph_format.space_after = Pt(8)
    r_t = p_t.add_run(titulo)
    r_t.font.size = Pt(24)
    r_t.font.bold = True
    r_t.font.color.rgb = BLANCO

    p_s = cell.add_paragraph()
    p_s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_s.paragraph_format.space_before = Pt(4)
    p_s.paragraph_format.space_after = Pt(24)
    r_s = p_s.add_run(subtitulo)
    r_s.font.size = Pt(13)
    r_s.font.color.rgb = DORADO

    if descripcion:
        doc.add_paragraph()
        p_d = doc.add_paragraph()
        p_d.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_d = p_d.add_run(descripcion)
        r_d.font.size = Pt(11)
        r_d.font.italic = True
        r_d.font.color.rgb = GRIS_TEXTO

    doc.add_paragraph()
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_m = p_meta.add_run("Sistema Comercial Operativo | León López Consultoría | Marzo 2026")
    r_m.font.size = Pt(10)
    r_m.font.color.rgb = GRIS_TEXTO


def h1(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = AZUL
    return p


def h2(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = DORADO
    return p


def h3(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = AZUL
    return p


def body(doc, text: str, italic: bool = False, indent: float = 0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    r = p.add_run(text)
    r.font.size = Pt(11)
    r.font.color.rgb = NEGRO
    r.font.italic = italic
    return p


def bullet(doc, text: str, level: int = 1, bold_prefix: str = ""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(level * 0.6)
    if bold_prefix:
        rb = p.add_run(bold_prefix + " ")
        rb.font.bold = True
        rb.font.color.rgb = AZUL
        rb.font.size = Pt(11)
    r = p.add_run(text)
    r.font.size = Pt(11)
    r.font.color.rgb = NEGRO
    return p


def numbered(doc, text: str, num: int, bold_prefix: str = ""):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(1.0)
    rn = p.add_run(f"{num}. ")
    rn.font.bold = True
    rn.font.color.rgb = DORADO
    rn.font.size = Pt(11)
    if bold_prefix:
        rb = p.add_run(bold_prefix + " ")
        rb.font.bold = True
        rb.font.color.rgb = AZUL
        rb.font.size = Pt(11)
    r = p.add_run(text)
    r.font.size = Pt(11)
    r.font.color.rgb = NEGRO
    return p


def section_box(doc, title: str, content_lines: list, bg_hex: str = "F5F5F5",
                title_color: RGBColor = None, icon: str = ""):
    """Caja con título coloreado y contenido."""
    if title_color is None:
        title_color = AZUL
    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.cell(0, 0)
    set_cell_bg(cell, bg_hex)

    p_t = cell.paragraphs[0]
    p_t.paragraph_format.space_before = Pt(8)
    p_t.paragraph_format.space_after = Pt(4)
    p_t.paragraph_format.left_indent = Cm(0.4)
    rt = p_t.add_run((icon + " " if icon else "") + title.upper())
    rt.font.bold = True
    rt.font.size = Pt(11)
    rt.font.color.rgb = title_color

    for line in content_lines:
        pl = cell.add_paragraph()
        pl.paragraph_format.left_indent = Cm(0.6)
        pl.paragraph_format.space_before = Pt(2)
        pl.paragraph_format.space_after = Pt(2)
        rl = pl.add_run(line)
        rl.font.size = Pt(10)
        rl.font.color.rgb = NEGRO

    p_end = cell.add_paragraph()
    p_end.paragraph_format.space_after = Pt(8)
    doc.add_paragraph()


def semaforo_table(doc, items: list):
    """Tabla con semáforo visual. items = [(label, estado, descripcion)] donde estado: verde/amarillo/rojo"""
    colores = {"verde": "27AE60", "amarillo": "F39C12", "rojo": "C0392B", "gris": "95A5A6"}
    tbl = doc.add_table(rows=1, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    # Header
    heads = ["DIMENSIÓN", "ESTADO", "OBSERVACIÓN"]
    head_widths = [Inches(2.5), Inches(1.2), Inches(2.8)]
    for i, h in enumerate(heads):
        cell = tbl.cell(0, i)
        cell.width = head_widths[i]
        set_cell_bg(cell, "1A2B4A")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = BLANCO

    for label, estado, desc in items:
        row = tbl.add_row()
        c0 = row.cells[0]
        c0.width = head_widths[0]
        p0 = c0.paragraphs[0]
        p0.paragraph_format.left_indent = Cm(0.2)
        r0 = p0.add_run(label)
        r0.font.size = Pt(10)
        r0.font.bold = True
        r0.font.color.rgb = AZUL

        c1 = row.cells[1]
        c1.width = head_widths[1]
        set_cell_bg(c1, colores.get(estado.lower(), "95A5A6"))
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p1.add_run(estado.upper())
        r1.font.size = Pt(10)
        r1.font.bold = True
        r1.font.color.rgb = BLANCO

        c2 = row.cells[2]
        c2.width = head_widths[2]
        p2 = c2.paragraphs[0]
        p2.paragraph_format.left_indent = Cm(0.2)
        r2 = p2.add_run(desc)
        r2.font.size = Pt(10)
        r2.font.color.rgb = NEGRO

    doc.add_paragraph()


def firma_block(doc, label: str = "Firma y fecha"):
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    r = p.add_run(f"{label}: ________________________________    Fecha: ________________")
    r.font.size = Pt(11)
    r.font.color.rgb = GRIS_TEXTO


def add_checklist(doc, items: list, title: str = ""):
    if title:
        h3(doc, title)
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Cm(0.5)
        r = p.add_run(f"☐  {item}")
        r.font.size = Pt(11)
        r.font.color.rgb = NEGRO
